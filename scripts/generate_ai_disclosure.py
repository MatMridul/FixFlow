"""Generate the official Samsung PRISM GenAI Hackathon 3.0 AI Disclosure Document.
Saves to SRMIST_FixFlow_AI_Disclosure.docx and LangAI3.0_AI_Disclosure.docx.
"""
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path

OUTPUT_PATH_1 = Path(r"C:\Mridul\Programs\FixFlow\SRMIST_FixFlow_AI_Disclosure.docx")
OUTPUT_PATH_2 = Path(r"C:\Mridul\Programs\FixFlow\LangAI3.0_AI_Disclosure.docx")

doc = Document()

# Set page margins
sections = doc.sections
for section in sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

# Title
title_p = doc.add_paragraph()
title_run = title_p.add_run("SAMSUNG PRISM | Gen AI Hackathon 3.0\nAI USAGE DISCLOSURE FORM")
title_run.font.name = "Arial"
title_run.font.size = Pt(18)
title_run.font.bold = True
title_run.font.color.rgb = RGBColor(26, 75, 140)
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph()

# 1. Team Details
h1 = doc.add_heading("1. Team Details", level=2)
h1.runs[0].font.color.rgb = RGBColor(26, 75, 140)

p = doc.add_paragraph()
p.add_run("Team Name: ").bold = True
p.add_run("SRMIST_FixFlow\n")
p.add_run("Project / Product Name: ").bold = True
p.add_run("FixFlow — Smart Guided Troubleshooting Engine (Theme 02)\n")
p.add_run("Organization / Institution: ").bold = True
p.add_run("SRM Institute of Science and Technology (SRMIST)\n")
p.add_run("Submission Date: ").bold = True
p.add_run("September 29, 2026\n")
p.add_run("GitHub Repository: ").bold = True
p.add_run("https://github.com/MatMridul/FixFlow\n")
p.add_run("Release Tag: ").bold = True
p.add_run("PRISM_GENAI_HACKATHON_Y2026")

# 2. AI Usage Declaration
h2 = doc.add_heading("2. AI Usage Declaration", level=2)
h2.runs[0].font.color.rgb = RGBColor(26, 75, 140)
p = doc.add_paragraph()
p.add_run("Did your team use any Artificial Intelligence (AI) in developing this project? ").bold = True
p.add_run("[X] Yes    [ ] No\n")
p.add_run("Declaration: ").bold = True
p.add_run("AI was leveraged responsibly for pair programming, test synthesis, and API integration, alongside custom mathematical algorithms (BM25, Graph Reranking, Compositional Gating).")

# 3. Purpose of AI Usage
h3 = doc.add_heading("3. Purpose of AI Usage (Brief Details)", level=2)
h3.runs[0].font.color.rgb = RGBColor(26, 75, 140)
purposes = [
    ("Idea generation / brainstorming", "Evaluated dual-process system patterns (Fast-Path Cache vs Slow-Path Extractor)."),
    ("Code generation or assistance", "Pair programming on Pydantic schemas, FastAPI route wiring, and graph traversal routines."),
    ("UI / UX design", "Assisted in One UI 6.1 styling rules, dark mode calibration, and WCAG accessibility standards."),
    ("Content creation", "Formatted technical Markdown documentation and API docstrings."),
    ("Data analysis", "Analyzed calibration set distributions, retrieval margins, and latency percentiles."),
    ("Testing / debugging", "Generated edge-case pytest fixtures and synthetic boundary validation tests."),
]
for cat, desc in purposes:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(f"{cat}: ").bold = True
    p.add_run(desc)

# 4. Feature Origin Classification
h4 = doc.add_heading("4. Feature Origin Classification", level=2)
h4.runs[0].font.color.rgb = RGBColor(26, 75, 140)

features = [
    {
        "name": "Feature 1: Compositional Gated Cache (Novelty N1 & N2)",
        "origin": "Both (Self-Architected + AI Assistance)",
        "desc": "Conceived and mathematically formalized multi-intent clause decomposition and intent-signature gating; AI assisted with SQLite WAL schema scaffolding and regex tokenization.",
    },
    {
        "name": "Feature 2: Settings Screen Graph Path Reranker (Novelty N3)",
        "origin": "Both (Self-Architected + AI Assistance)",
        "desc": "Designed directed graph representation of Samsung One UI Settings menus; AI helped implement Dijkstra shortest-path traversal and parent-menu deduplication.",
    },
    {
        "name": "Feature 3: Closed-Loop Verification Deeplinks (Novelty N4)",
        "origin": "Self-Generated",
        "desc": "Devised closed-loop execution architecture binding validationDeeplink to examine toggle states and auto-skip satisfied steps.",
    },
    {
        "name": "Feature 4: SIIS Knowledge Auto-Retriever",
        "origin": "Both (Self-Architected + AI Assistance)",
        "desc": "Engineered hybrid BM25 and TF-IDF cosine similarity retrieval over Samsung internal SIIS corpus to support freeform user queries without manual article pasting.",
    },
    {
        "name": "Feature 5: One UI 6.1 Interactive Phone Simulation",
        "origin": "Both (Self-Architected + AI Assistance)",
        "desc": "Built vanilla JS/CSS Galaxy S24 interactive device emulator; AI assisted with CSS keyframe timing and SVG battery/clock styling.",
    },
]

for f in features:
    p = doc.add_paragraph()
    p.add_run(f["name"] + "\n").bold = True
    p.add_run("Classification: ").bold = True
    p.add_run(f["origin"] + "\n")
    p.add_run("Description: ").bold = True
    p.add_run(f["desc"])

# 5. Ethical & Compliance Confirmation
h5 = doc.add_heading("5. Ethical & Compliance Confirmation", level=2)
h5.runs[0].font.color.rgb = RGBColor(26, 75, 140)
p = doc.add_paragraph()
p.add_run("AI usage complies with Samsung PRISM guidelines and policies: ").bold = True
p.add_run("YES [X]\n")
p.add_run("No proprietary or copyrighted data misused: ").bold = True
p.add_run("I AGREE [X]\n")
p.add_run("API Keys and Secrets: ").bold = True
p.add_run("All credentials managed strictly via gitignored environment variables (.env); zero hardcoded keys in repository.")

# 6. Declaration & Sign-Off
h6 = doc.add_heading("6. Declaration & Sign-Off", level=2)
h6.runs[0].font.color.rgb = RGBColor(26, 75, 140)
p = doc.add_paragraph()
p.add_run("Name of Team Representative: ").bold = True
p.add_run("Mridul Mathur\n")
p.add_run("Role: ").bold = True
p.add_run("Team Lead & Full-Stack AI Engineer\n")
p.add_run("Institution: ").bold = True
p.add_run("SRM Institute of Science and Technology\n")
p.add_run("Date: ").bold = True
p.add_run("September 29, 2026\n")
p.add_run("Signature: ").bold = True
p.add_run("Mridul Mathur")

doc.save(str(OUTPUT_PATH_1))
doc.save(str(OUTPUT_PATH_2))
print("SUCCESS: Generated AI Disclosure forms at:")
print(f"  1. {OUTPUT_PATH_1}")
print(f"  2. {OUTPUT_PATH_2}")
